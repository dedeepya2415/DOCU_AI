from pathlib import Path

import fitz
import easyocr
import cv2
import numpy as np
from fastapi import HTTPException

# Singleton: initialize once at module load to avoid reloading models per request
_reader = easyocr.Reader(['en'], gpu=False)


class OCRService:

    def __init__(self):
        self.reader = _reader

    def extract_text(self, file_path: str):

        extension = Path(file_path).suffix.lower()

        try:
            if extension == ".pdf":
                return self.extract_from_pdf(file_path)
            elif extension == ".docx":
                return self.extract_from_docx(file_path)
            return self.extract_from_image(file_path)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(
                status_code=400,
                detail=f"Failed to read file '{Path(file_path).name}': {str(e)}"
            )

    def extract_from_pdf(self, pdf_path: str):

        document = fitz.open(pdf_path)

        text = ""

        for page in document:
            page_text = page.get_text()

            if page_text.strip():
                text += page_text + "\n"
            else:
                pix = page.get_pixmap(dpi=200)
                img_array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                
                if pix.n == 4:
                    img_array = cv2.cvtColor(img_array, cv2.COLOR_RGBA2RGB)
                elif pix.n == 1:
                    img_array = cv2.cvtColor(img_array, cv2.COLOR_GRAY2RGB)
                    
                gray = cv2.cvtColor(img_array, cv2.COLOR_RGB2GRAY)
                result = self.reader.readtext(gray, detail=0)
                text += "\n".join(result) + "\n"

        document.close()

        return text

    def extract_from_image(self, image_path: str):

        img = cv2.imread(image_path)
        if img is None:
            result = self.reader.readtext(image_path, detail=0)
            return "\n".join(result)
            
        # Resize image to a maximum dimension to drastically speed up OCR
        max_dim = 1024
        h, w = img.shape[:2]
        if max(h, w) > max_dim:
            scale = max_dim / max(h, w)
            img = cv2.resize(img, (int(w * scale), int(h * scale)))

        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        
        # Removed fastNlMeansDenoising as it is extremely slow on CPU.
        # EasyOCR is usually robust enough to handle basic image noise.
        result = self.reader.readtext(gray, detail=0)

        return "\n".join(result)

    def extract_from_docx(self, docx_path: str):
        from docx import Document
        doc = Document(docx_path)
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            text.append(p.text)
        return "\n".join(text)