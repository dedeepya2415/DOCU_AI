import docx

def create_master_template():
    doc = docx.Document()
    doc.add_heading('Sale Deed', 0)
    doc.add_paragraph('This Sale Deed is made and executed on {{ registration_date }}.')
    doc.add_heading('BETWEEN', level=2)
    doc.add_paragraph('{{ seller_name }}, S/o {{ seller_father_name }}, aged about 63 years, bearing Aadhaar No. {{ seller_aadhaar }} and PAN No. {{ seller_pan }}, residing at {{ seller_address }}.')
    doc.add_paragraph('(hereinafter referred to as the "VENDOR")')
    doc.add_heading('AND', level=2)
    doc.add_paragraph('{{ buyer_name }}, bearing Aadhaar No. {{ buyer_aadhaar }} and PAN No. {{ buyer_pan }}, residing at {{ buyer_address }}.')
    doc.add_paragraph('(hereinafter referred to as the "PURCHASER")')
    doc.add_heading('WHEREAS', level=2)
    doc.add_paragraph('The VENDOR is the absolute owner of the property located at {{ property_address }}, containing an area of {{ area }} bearing survey number {{ survey_number }}.')
    doc.add_paragraph('The VENDOR has agreed to sell the Schedule Property to the PURCHASER for a total sale consideration of {{ sale_price }}.')
    
    doc.save(r"C:\Users\datla\OneDrive\Desktop\Master_Sale_Deed.docx")
    print("Master template created!")

if __name__ == "__main__":
    create_master_template()
